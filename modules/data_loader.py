"""
Modul Pemuat Data & Export File
Terminal Teluk Lamong - Pelindo

Menyediakan fungsi pembacaan streaming file Excel (.xlsx) dengan pelaporan progres riil,
hemat memori (read_only mode openpyxl), serta generator file Excel hasil ekspor.
"""

from io import BytesIO
import time
import openpyxl
from openpyxl.styles import Font
import pandas as pd
import streamlit as st
from modules.ui import format_number


def baca_file(file_bytes: bytes, filename: str, progress_callback=None) -> dict[str, pd.DataFrame]:
    """
    Membaca file data operasional (.xlsx) dengan pelaporan progres riil
    dan efisiensi memori tingkat tinggi (streaming SAX parser via openpyxl).
    Dikhususkan hanya untuk file Excel (.xlsx) guna menjamin validitas tipe data
    dan akurasi kalkulasi timestamp.
    """
    fname_lower = filename.lower()

    # Validasi format file: hanya menerima .xlsx
    if not fname_lower.endswith(".xlsx"):
        raise ValueError(
            f"Format file '{filename}' tidak didukung. "
            "Aplikasi ini dikhususkan hanya untuk file Excel (.xlsx) guna menjamin "
            "akurasi kalkulasi waktu dan integritas tipe data operasional."
        )

    # ------------------------------------------------------------
    # Penanganan File Excel Modern (.xlsx) via openpyxl streaming
    # ------------------------------------------------------------
    if progress_callback:
        progress_callback(5, "Menganalisis struktur workbook Excel (.xlsx)...", "Membuka lembar kerja...")

    bio = BytesIO(file_bytes)
    wb = openpyxl.load_workbook(bio, read_only=True, data_only=True)
    sheet_names = wb.sheetnames
    num_sheets = len(sheet_names)

    sheets_dict = {}
    last_update = 0.0

    for s_idx, sname in enumerate(sheet_names):
        ws = wb[sname]
        max_row = ws.max_row  # Mendapatkan total baris lembar kerja
        base_pct = int(8 + (s_idx / num_sheets) * 88)
        sheet_weight = 88.0 / num_sheets

        if progress_callback:
            target_info = f"{format_number(max_row)} total baris" if max_row else "memulai baris..."
            progress_callback(
                base_pct,
                f"Mengekstrak sheet '{sname}' ({s_idx + 1}/{num_sheets})...",
                target_info,
            )

        rows_iter = ws.iter_rows(values_only=True)
        try:
            header = next(rows_iter)
        except StopIteration:
            sheets_dict[sname] = pd.DataFrame()
            continue

        # Bersihkan nama kolom header
        clean_headers = [
            str(c).strip() if c is not None and str(c).strip() != "" else f"Col_{i}"
            for i, c in enumerate(header)
        ]

        rows = []
        last_pct = base_pct
        for r_idx, row in enumerate(rows_iter):
            rows.append(row)
            now = time.time()
            # Batasi frekuensi callback agar stabil, mulus, dan tidak flickering
            if now - last_update > 0.20:
                if max_row and max_row > 1:
                    sheet_prog = min(1.0, (r_idx + 1) / (max_row - 1))
                else:
                    sheet_prog = min(0.95, (r_idx + 1) / 100000)

                current_pct = int(base_pct + sheet_prog * sheet_weight)
                current_pct = min(96, max(base_pct, current_pct))

                if current_pct > last_pct:
                    last_update = now
                    last_pct = current_pct
                    if progress_callback:
                        row_info = f"Baris {format_number(r_idx + 1)}" + (f" / {format_number(max_row)}" if max_row else "")
                        progress_callback(
                            current_pct,
                            f"Mengekstrak '{sname}'...",
                            row_info,
                        )

        df = pd.DataFrame(rows, columns=clean_headers)
        sheets_dict[sname] = df

    wb.close()

    if progress_callback:
        total_all_rows = sum(len(df) for df in sheets_dict.values())
        progress_callback(
            100,
            "Data siap dianalisis!",
            f"{format_number(total_all_rows)} baris siap",
        )
    return sheets_dict


def build_excel_data_only(out_df: pd.DataFrame) -> bytes:
    """
    Membangun file Excel hasil analisis (.xlsx) ringan dan cepat (hanya sheet Data),
    menggunakan openpyxl dengan header tebal dan freeze pane.
    """
    bio = BytesIO()
    with pd.ExcelWriter(bio, engine="openpyxl") as writer:
        out_df.to_excel(writer, index=False, sheet_name="Data")
        ws = writer.sheets["Data"]
        for cell in ws[1]:
            cell.font = Font(bold=True)
        ws.freeze_panes = "A2"
    bio.seek(0)
    return bio.getvalue()
